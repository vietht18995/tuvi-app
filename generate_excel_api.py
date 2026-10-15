import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_for_chart(tuvi_data, output_excel_path):
    wb = openpyxl.Workbook()

    # Setup Sheet 1: Lá số Tử Vi (Trực quan)
    ws1 = wb.active
    ws1.title = "Lá số Tử Vi (Trực quan)"
    ws1.views.sheetView[0].showGridLines = True

    # Setup Sheet 2: Chi tiết 12 Cung (Dạng bảng)
    ws2 = wb.create_sheet(title="Chi tiết 12 Cung (Dạng bảng)")
    ws2.views.sheetView[0].showGridLines = True

    # Define styles and colors
    colors_hex = {
        "Mộc": "E2EFDA",     # Light green
        "Hỏa": "FCE4D6",     # Light red/orange
        "Thổ": "FFF2CC",     # Light yellow
        "Kim": "F2F2F2",     # Light grey
        "Thủy": "DDEBF7",    # Light blue
        "Center": "FFFDF0",   # Soft cream color for center panel
        "White": "FFFFFF",
        "GrayBorder": "D9D9D9"
    }

    font_colors_hex = {
        "Mộc": "375623",     # Dark green
        "Hỏa": "C65911",     # Dark orange/red
        "Thổ": "7F6000",     # Dark golden-brown
        "Kim": "595959",     # Dark grey
        "Thủy": "1F4E78",    # Dark blue
        "Positive": "1F4E78", # Dark blue for positive stars
        "Negative": "C00000", # Red for negative/hung stars
        "Neutral": "000000"
    }

    # Border styles
    thin_border_side = Side(style='thin', color='D9D9D9')
    medium_border_side = Side(style='medium', color='595959')

    # Borders for the 8x4 blocks
    def apply_chamber_borders_and_fill(ws, start_row, start_col, fill_color):
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for r in range(start_row, start_row + 8):
            for c in range(start_col, start_col + 4):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill
                
                # Apply border boundaries
                t = medium_border_side if r == start_row else None
                b = medium_border_side if r == start_row + 7 else None
                l = medium_border_side if c == start_col else None
                rg = medium_border_side if c == start_col + 3 else None
                
                # Inner thin borders
                if r == start_row:
                    b = thin_border_side
                elif r == start_row + 1:
                    b = thin_border_side
                elif r == start_row + 5:
                    b = thin_border_side
                elif r == start_row + 6:
                    b = thin_border_side
                    
                if c == start_col + 1 and r >= start_row + 2 and r <= start_row + 5:
                    rg = thin_border_side
                    
                cell.border = Border(top=t, bottom=b, left=l, right=rg)

    def write_chamber(ws, grid_r, grid_c, chamber):
        start_row = 3 + grid_r * 8
        start_col = 2 + grid_c * 4
        
        fill_color = colors_hex[chamber["element"]]
        apply_chamber_borders_and_fill(ws, start_row, start_col, fill_color)
        
        font_bold = Font(name="Calibri", size=11, bold=True, color="000000")
        font_regular = Font(name="Calibri", size=9, color="333333")
        font_italic = Font(name="Calibri", size=9, italic=True, color="555555")
        
        # Row 0: Can Chi + Element & Chamber Name
        cell_info = ws.cell(row=start_row, column=start_col)
        cell_info.value = f"{chamber['can']}.{chamber['chi']} ({chamber['polarity']}{chamber['element']})"
        cell_info.font = Font(name="Calibri", size=9, bold=True, color=font_colors_hex[chamber["element"]])
        cell_info.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + 1)
        
        cell_name = ws.cell(row=start_row, column=start_col + 2)
        name_suffix = " <THÂN>" if chamber["is_than"] else ""
        cell_name.value = f"{chamber['name']}{name_suffix}"
        cell_name.font = Font(name="Calibri", size=12, bold=True, color="A52A2A" if chamber["is_than"] else "000000")
        cell_name.alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=start_row, start_column=start_col + 2, end_row=start_row, end_column=start_col + 3)
        
        # Row 1: Đại Hạn & Nguyệt Hạn
        cell_dh = ws.cell(row=start_row + 1, column=start_col)
        cell_dh.value = f"Đại hạn: {chamber['dai_han']}"
        cell_dh.font = font_regular
        cell_dh.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=start_row + 1, start_column=start_col, end_row=start_row + 1, end_column=start_col + 1)
        
        cell_nh = ws.cell(row=start_row + 1, column=start_col + 2)
        cell_nh.value = f"Nguyệt hạn: Th.{chamber['nguyet_han']}"
        cell_nh.font = font_italic
        cell_nh.alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=start_row + 1, start_column=start_col + 2, end_row=start_row + 1, end_column=start_col + 3)
        
        # Row 2-5: Main Stars
        left_lines = []
        for ct in chamber["chinh_tinh"]:
            status_str = f"({ct['status']})" if ct.get("status") else ""
            left_lines.append(f"★ {ct['name']} {status_str}")
        if left_lines and chamber["cat_tinh"]:
            left_lines.append("----------------")
        for cat in chamber["cat_tinh"]:
            luu_prefix = "L." if cat.get("is_luu") else ""
            left_lines.append(f"• {luu_prefix}{cat['name']}")
            
        cell_left = ws.cell(row=start_row + 2, column=start_col)
        cell_left.value = "\n".join(left_lines)
        cell_left.font = Font(name="Calibri", size=9, bold=True, color=font_colors_hex["Positive"])
        cell_left.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.merge_cells(start_row=start_row + 2, start_column=start_col, end_row=start_row + 5, end_column=start_col + 1)
        
        right_lines = []
        for hung in chamber["hung_tinh"]:
            status_str = f"({hung['status']})" if hung.get("status") else ""
            luu_prefix = "L." if hung.get("is_luu") else ""
            right_lines.append(f"x {luu_prefix}{hung['name']} {status_str}")
            
        cell_right = ws.cell(row=start_row + 2, column=start_col + 2)
        cell_right.value = "\n".join(right_lines)
        cell_right.font = Font(name="Calibri", size=9, color=font_colors_hex["Negative"])
        cell_right.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.merge_cells(start_row=start_row + 2, start_column=start_col + 2, end_row=start_row + 5, end_column=start_col + 3)
        
        # Row 6: Vòng Tràng Sinh
        cell_ts = ws.cell(row=start_row + 6, column=start_col)
        cell_ts.value = f"Tràng sinh: {chamber['trang_sinh']}"
        cell_ts.font = Font(name="Calibri", size=9, bold=True, color="555555")
        cell_ts.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=start_row + 6, start_column=start_col, end_row=start_row + 6, end_column=start_col + 3)
        
        # Row 7: Tuần / Triệt | Đại vận | Lưu Niên
        tt_str = chamber["tuan_triet"] if chamber["tuan_triet"] else ""
        dv_ln_str = f"ĐV.{chamber['dai_van_luu_nien']['dv']} | LN.{chamber['dai_van_luu_nien']['ln']}"
        combined_bottom = f"{tt_str} | {dv_ln_str}" if tt_str else dv_ln_str
        
        cell_bottom = ws.cell(row=start_row + 7, column=start_col)
        cell_bottom.value = combined_bottom
        if tt_str:
            cell_bottom.font = Font(name="Calibri", size=9, bold=True, color="A52A2A")
        else:
            cell_bottom.font = font_regular
        cell_bottom.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=start_row + 7, start_column=start_col, end_row=start_row + 7, end_column=start_col + 3)

    grid_mapping = {
        "dien_trach": (0, 0), "quan_loc": (0, 1), "no_boc": (0, 2), "thien_di": (0, 3),
        "phuc_duc": (1, 0),                                         "tat_ach": (1, 3),
        "phu_mau": (2, 0),                                           "tai_bach": (2, 3),
        "menh": (3, 0),       "huynh_de": (3, 1), "phu_the": (3, 2), "tu_tuc": (3, 3)
    }

    chambers_dict = {c["id"]: c for c in tuvi_data["chambers"]}

    for c_id, coord in grid_mapping.items():
        if c_id in chambers_dict:
            write_chamber(ws1, coord[0], coord[1], chambers_dict[c_id])

    # Write Central Panel (Personal Info)
    fill_center = PatternFill(start_color=colors_hex["Center"], end_color=colors_hex["Center"], fill_type="solid")
    for r in range(11, 27):
        for c in range(6, 14):
            cell = ws1.cell(row=r, column=c)
            cell.fill = fill_center
            
            t = medium_border_side if r == 11 else None
            b = medium_border_side if r == 26 else None
            l = medium_border_side if c == 6 else None
            rg = medium_border_side if c == 13 else None
            cell.border = Border(top=t, bottom=b, left=l, right=rg)

    personal = tuvi_data["personal_info"]

    # Title
    ws1.merge_cells("F12:M13")
    title_cell = ws1["F12"]
    title_cell.value = "LÁ SỐ TỬ VI"
    title_cell.font = Font(name="Calibri", size=18, bold=True, color="800000")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Subtitle
    ws1.merge_cells("F14:M14")
    sub_cell = ws1["F14"]
    sub_cell.value = "TỰ ĐỘNG KHỞI TẠO - TRA CỨU CÁ NHÂN"
    sub_cell.font = Font(name="Calibri", size=9, italic=True, color="555555")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    details = [
        ("Họ và tên:", personal["name"], "Năm sinh:", f"{personal['birth_year_solar']} ({personal['birth_year_lunar']})"),
        ("Tháng sinh (DL/AL):", f"{personal['birth_month_solar']} ({personal['birth_month_lunar']})", "Thiên can/Địa chi:", personal["birth_month_lunar_stem"]),
        ("Ngày sinh (DL/AL):", f"{personal['birth_day_solar']} ({personal['birth_day_lunar']})", "Thiên can/Địa chi:", personal["birth_day_lunar_stem"]),
        ("Giờ sinh:", personal["birth_hour"], "Thiên can/Địa chi:", personal["birth_hour_stem"]),
        ("Bản mệnh:", personal["element_fate"], "Cục:", personal["element_cuc"]),
        ("Mối quan hệ:", personal["relation_fate_cuc"], "Cân lượng:", personal["weight"]),
        ("Chủ mệnh:", personal["fate_lord"], "Chủ thân:", personal["body_lord"]),
        ("Lai nhân cung:", personal["lai_nhan_cung"], "Giới tính / Loại:", f"{personal['gender']} (Âm Dương)")
    ]

    for idx, (label1, val1, label2, val2) in enumerate(details):
        r = 16 + idx
        ws1.cell(row=r, column=6, value=label1).font = Font(name="Calibri", size=9, bold=True, color="333333")
        ws1.cell(row=r, column=6).alignment = Alignment(horizontal="left", vertical="center")
        
        ws1.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
        val1_cell = ws1.cell(row=r, column=7, value=val1)
        val1_cell.font = Font(name="Calibri", size=9, color="000000")
        val1_cell.alignment = Alignment(horizontal="left", vertical="center")
        
        ws1.cell(row=r, column=10, value=label2).font = Font(name="Calibri", size=9, bold=True, color="333333")
        ws1.cell(row=r, column=10).alignment = Alignment(horizontal="left", vertical="center")
        
        ws1.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
        val2_cell = ws1.cell(row=r, column=11, value=val2)
        val2_cell.font = Font(name="Calibri", size=9, color="000000")
        val2_cell.alignment = Alignment(horizontal="left", vertical="center")

    # Sheet Title
    ws1.merge_cells("B1:Q1")
    sheet_title = ws1["B1"]
    sheet_title.value = f"LÁ SỐ TỬ VI - {personal['name'].upper()} (NĂM XEM: {personal['analysis_year_lunar']} - {personal['analysis_year_solar']})"
    sheet_title.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    sheet_title.fill = PatternFill(start_color="800000", end_color="800000", fill_type="solid")
    sheet_title.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 40

    # Add Legend
    ws1.merge_cells("B36:Q36")
    legend_cell = ws1["B36"]
    legend_cell.value = "CHÚ GIẢI NGŨ HÀNH CỦA CUNG & SAO:   " \
                        "[■ MỘC] Xanh lá (Mộc)   " \
                        "[■ HỎA] Cam/Đỏ (Hỏa)   " \
                        "[■ THỔ] Vàng (Thổ)   " \
                        "[■ KIM] Xám (Kim)   " \
                        "[■ THỦY] Xanh dương (Thủy)       " \
                        "(*) Ký hiệu độ miếu hãm: (M): Miếu địa, (V): Vượng địa, (Đ): Đắc địa, (B): Bình hòa, (H): Hãm địa"
    legend_cell.font = Font(name="Calibri", size=9, italic=True, color="555555")
    legend_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[36].height = 25

    # Adjust column widths for Sheet 1
    for col_idx in range(2, 18):
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = 11

    # Adjust row heights for Sheet 1
    for row_idx in range(3, 35):
        ws1.row_dimensions[row_idx].height = 20
    ws1.row_dimensions[2].height = 10


    # ----------------------------------------------------
    # Setup Sheet 2: Chi tiết 12 Cung (Dạng bảng)
    headers2 = [
        "STT", "Tên Cung", "Địa Chi", "Hành Cung", "Âm Dương", 
        "Đại Hạn", "Nguyệt Hạn", "Tràng Sinh", "Chính Tinh", 
        "Cát Tinh", "Hung Tinh / Sát Tinh", "Tuần / Triệt", "Đại Vận", "Lưu Niên"
    ]

    ws2.row_dimensions[1].height = 28
    for c_idx, h_text in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=c_idx, value=h_text)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium_border_side, bottom=medium_border_side, left=thin_border_side, right=thin_border_side)

    # Write tabular data
    clockwise_order = ["menh", "phu_mau", "phuc_duc", "dien_trach", "quan_loc", "no_boc", "thien_di", "tat_ach", "tai_bach", "tu_tuc", "phu_the", "huynh_de"]
    
    # Filter only existing chambers in dict
    chambers_list = [chambers_dict[cid] for cid in clockwise_order if cid in chambers_dict]

    for r_idx, chamber in enumerate(chambers_list, 2):
        ws2.row_dimensions[r_idx].height = 32
        
        ws2.cell(row=r_idx, column=1, value=r_idx - 1)
        
        name_suffix = " (Thân)" if chamber["is_than"] else ""
        ws2.cell(row=r_idx, column=2, value=f"{chamber['name']}{name_suffix}")
        ws2.cell(row=r_idx, column=3, value=f"{chamber['can']}.{chamber['chi']}")
        ws2.cell(row=r_idx, column=4, value=chamber["element"])
        ws2.cell(row=r_idx, column=5, value="Dương" if chamber["polarity"] == "+" else "Âm")
        ws2.cell(row=r_idx, column=6, value=chamber["dai_han"])
        ws2.cell(row=r_idx, column=7, value=f"Th.{chamber['nguyet_han']}")
        ws2.cell(row=r_idx, column=8, value=chamber["trang_sinh"])
        
        ct_strs = []
        for ct in chamber["chinh_tinh"]:
            status_str = f"({ct['status']})" if ct.get("status") else ""
            ct_strs.append(f"{ct['name']} {status_str}")
        ws2.cell(row=r_idx, column=9, value=", ".join(ct_strs))
        
        cat_strs = []
        for cat in chamber["cat_tinh"]:
            luu_prefix = "L." if cat.get("is_luu") else ""
            cat_strs.append(f"{luu_prefix}{cat['name']}")
        ws2.cell(row=r_idx, column=10, value=", ".join(cat_strs))
        
        hung_strs = []
        for hung in chamber["hung_tinh"]:
            status_str = f"({hung['status']})" if hung.get("status") else ""
            luu_prefix = "L." if hung.get("is_luu") else ""
            hung_strs.append(f"{luu_prefix}{hung['name']}{status_str}")
        ws2.cell(row=r_idx, column=11, value=", ".join(hung_strs))
        
        ws2.cell(row=r_idx, column=12, value=chamber["tuan_triet"] if chamber["tuan_triet"] else "Không")
        ws2.cell(row=r_idx, column=13, value=chamber["dai_van_luu_nien"]["dv"])
        ws2.cell(row=r_idx, column=14, value=chamber["dai_van_luu_nien"]["ln"])
        
        for col_idx in range(1, 15):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(top=thin_border_side, bottom=thin_border_side, left=thin_border_side, right=thin_border_side)
            
            if col_idx == 4:
                elem = cell.value
                cell.fill = PatternFill(start_color=colors_hex[elem], end_color=colors_hex[elem], fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color=font_colors_hex[elem])
            elif col_idx == 9:
                cell.font = Font(name="Calibri", size=10, bold=True, color="000000")
            elif col_idx == 10:
                cell.font = Font(name="Calibri", size=9, color=font_colors_hex["Positive"])
            elif col_idx == 11:
                cell.font = Font(name="Calibri", size=9, color=font_colors_hex["Negative"])
            elif col_idx == 12:
                if cell.value != "Không":
                    cell.font = Font(name="Calibri", size=10, bold=True, color="A52A2A")
            else:
                cell.font = Font(name="Calibri", size=10)

    # Autofit columns
    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 10)

    # ----------------------------------------------------
    # Setup Sheet 3: Tứ Trụ - Bát Tự
    ws3 = wb.create_sheet(title="Tứ Trụ - Bát Tự")
    ws3.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws3.merge_cells("A1:I1")
    title3 = ws3["A1"]
    title3.value = f"BẢNG PHÂN TÍCH TỨ TRỤ - BÁT TỰ: {personal['name'].upper()}"
    title3.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title3.fill = PatternFill(start_color="800000", end_color="800000", fill_type="solid")
    title3.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 35
    
    # Bazi table headers
    bazi_headers = ["Thông tin", "Trụ Năm", "Trụ Tháng", "Trụ Ngày (Nhật Chủ)", "Trụ Giờ", "Trụ Phút", "Thai Cung", "Mệnh Cung", "Thai Tức"]
    ws3.row_dimensions[3].height = 24
    for c_idx, h_text in enumerate(bazi_headers, 1):
        cell = ws3.cell(row=3, column=c_idx, value=h_text)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=medium_border_side, bottom=medium_border_side, left=thin_border_side, right=thin_border_side)
        
    bazi_data = tuvi_data.get("bazi", {})
    pillars_data = bazi_data.get("pillars", {})
    
    rows_def = [
        ("Thiên Can", lambda p: p.get("can"), lambda p: p.get("can_elem")),
        ("Âm Dương Can", lambda p: p.get("can_polarity"), None),
        ("Thập Thần Can", lambda p: p.get("can_ten_god"), None),
        ("Địa Chi", lambda p: p.get("chi"), lambda p: p.get("chi_elem")),
        ("Âm Dương Chi", lambda p: p.get("chi_polarity"), None),
        ("Tàng Can (Bản khí)", lambda p: p.get("hidden_can"), lambda p: p.get("hidden_can_elem")),
        ("Thập Thần Tàng Can", lambda p: p.get("hidden_ten_god"), None)
    ]
    
    pillar_keys = ["year", "month", "day", "hour", "minute", "thai_cung", "menh_cung", "thai_tuc"]
    
    for r_idx, (label, val_func, elem_func) in enumerate(rows_def, 4):
        ws3.row_dimensions[r_idx].height = 24
        
        # Label cell
        cell_lbl = ws3.cell(row=r_idx, column=1, value=label)
        cell_lbl.font = Font(name="Calibri", size=10, bold=True)
        cell_lbl.alignment = Alignment(horizontal="left", vertical="center")
        cell_lbl.border = Border(top=thin_border_side, bottom=thin_border_side, left=thin_border_side, right=thin_border_side)
        cell_lbl.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
        # Pillar values
        for c_idx, key in enumerate(pillar_keys, 2):
            p = pillars_data.get(key, {})
            val = val_func(p) if p else ""
            cell_val = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell_val.alignment = Alignment(horizontal="center", vertical="center")
            cell_val.border = Border(top=thin_border_side, bottom=thin_border_side, left=thin_border_side, right=thin_border_side)
            
            # Element coloring
            if elem_func:
                elem = elem_func(p) if p else ""
                if elem in colors_hex:
                    cell_val.fill = PatternFill(start_color=colors_hex[elem], end_color=colors_hex[elem], fill_type="solid")
                    cell_val.font = Font(name="Calibri", size=10, bold=True, color=font_colors_hex[elem])
                else:
                    cell_val.font = Font(name="Calibri", size=10)
            else:
                cell_val.font = Font(name="Calibri", size=10)
                
            # Special formatting for Ten Gods
            if "Thập Thần" in label:
                if val == "Nhật Chủ (Bản Thân)":
                    cell_val.font = Font(name="Calibri", size=10, bold=True, color="800000")
                elif val in ["Chính Tài", "Thiên Tài", "Chính Quan", "Chính Ấn", "Thiên Ấn", "Thực Thần"]:
                    cell_val.font = Font(name="Calibri", size=10, bold=True, color="375623")
                elif val in ["Thương Quan", "Thất Sát", "Kiếp Tài"]:
                    cell_val.font = Font(name="Calibri", size=10, bold=True, color="C00000")
                    
    # Vuong Suy Box
    ws3.merge_cells("A13:G15")
    cell_vs = ws3["A13"]
    cell_vs.value = f"LUẬN GIẢI BÁT TỰ - NHẬT CHỦ VƯỢNG SUY:\n" \
                    f"• Nhật Chủ (Bản thân): {bazi_data.get('day_master')} ({bazi_data.get('day_master_polarity')}{bazi_data.get('day_master_elem')})\n" \
                    f"• Trạng thái Vượng Suy: {bazi_data.get('vuong_suy')}\n" \
                    f"• Luận giải ngắn gọn: Đương số có Nhật Chủ là {bazi_data.get('day_master')}. Mùa sinh (Địa chi tháng) ảnh hưởng mạnh mẽ nhất đến năng lượng của Nhật Chủ, giúp xác định mức độ vượng suy và lựa chọn Dụng Thần phù hợp để điều hòa bản mệnh."
    cell_vs.font = Font(name="Calibri", size=10)
    cell_vs.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    # Border for analysis block
    for r in range(13, 16):
        for c in range(1, 8):
            cell = ws3.cell(row=r, column=c)
            cell.fill = PatternFill(start_color="FFFDF0", end_color="FFFDF0", fill_type="solid")
            cell.border = Border(
                top=medium_border_side if r == 13 else None,
                bottom=medium_border_side if r == 15 else None,
                left=medium_border_side if c == 1 else None,
                right=medium_border_side if c == 7 else None
            )
            
    # Auto-fit ws3 columns
    for col in ws3.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 15)

    # Save
    wb.save(output_excel_path)
    print(f"Excel generated successfully at {output_excel_path}")
